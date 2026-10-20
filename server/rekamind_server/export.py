"""One document model of a meeting, three renderers: Markdown, Word (.docx) and JSON.

Everything a person would want to keep or hand over: who was there, the summary, decisions, the timeline,
the tasks (confirmed ones clearly apart from suggestions) and, optionally, the full transcript. Built on the
server so the web UI, the desktop app and scripts all get the same file, and so the access rules that guard
every other route guard the export too.
"""
from __future__ import annotations

import io
import json
import re
import unicodedata
from datetime import UTC, date, datetime
from typing import Literal
from zoneinfo import ZoneInfo

from .models import Meeting, Participant, Task
from .tz import DEFAULT as DEFAULT_TZ
from .tz import label, to_zone, zone

SCHEMA = "rekamind.meeting/v1"
TASKS_SCHEMA = "rekamind.tasks/v1"
TaskFilter = Literal["confirmed", "all"]
PRIORITY = {"low": "Rendah", "medium": "Sedang", "high": "Tinggi"}
STATUS = {"suggested": "Disarankan (belum ditinjau)", "confirmed": "Dikonfirmasi", "rejected": "Ditolak"}

# Code points XML 1.0 cannot carry: C0 controls except tab/LF/CR, lone surrogates, U+FFFE and U+FFFF.
_XML_BAD = dict.fromkeys((*range(0x00, 0x09), 0x0B, 0x0C, *range(0x0E, 0x20), *range(0xD800, 0xE000), 0xFFFE, 0xFFFF))


def clean(text: str | None) -> str:
    """Control characters (a stray byte from ASR or the LLM) make python-docx raise; drop them."""
    return (text or "").translate(_XML_BAD)


def clock(ms: int | None) -> str:
    if ms is None:
        return ""
    s = ms // 1000
    return f"{s // 3600}:{s % 3600 // 60:02d}:{s % 60:02d}" if s >= 3600 else f"{s // 60:02d}:{s % 60:02d}"


def minutes_text(m: int | None) -> str:
    if not m:
        return ""
    h, r = divmod(m, 60)
    return f"{h} j {r} mnt" if h and r else (f"{h} jam" if h else f"{r} menit")


def _who(t: Task, people: dict[str, Participant]) -> str | None:
    p = people.get(t.assignee_participant_id or "")
    return (p.name or p.label) if p else t.assignee_name


def task_dict(t: Task, people: dict[str, Participant]) -> dict:
    return {
        "id": t.id, "title": clean(t.title), "description": clean(t.description), "assignee": clean(_who(t, people)) or None,
        "deadline": t.deadline.isoformat() if t.deadline else None, "deadline_text": clean(t.deadline_text) or None,
        "deadline_is_approximate": bool(t.deadline_approximate), "estimate_minutes": t.estimate_minutes,
        "priority": t.priority, "status": t.status, "confidence": round(float(t.confidence), 2),
        "evidence": {"start_ms": t.evidence_start_ms, "quote": clean(t.evidence_quote)},
    }


def build_document(m: Meeting, *, transcript: bool = True, tasks: TaskFilter = "confirmed", tz: ZoneInfo | None = None) -> dict:
    tz = tz or zone(DEFAULT_TZ)
    people = {p.id: p for p in m.participants}
    chosen = [t for t in m.tasks if (t.status == "confirmed" if tasks == "confirmed" else t.status != "rejected")]
    order = {"confirmed": 0, "suggested": 1}
    chosen.sort(key=lambda t: (order.get(t.status, 2), t.created_at))
    doc: dict = {
        "schema": SCHEMA,
        "meeting": {"id": m.id, "title": clean(m.title), "started_at": m.started_at.isoformat() if m.started_at else None,
                    "ended_at": m.ended_at.isoformat() if m.ended_at else None, "language": m.language, "status": m.status,
                    "timezone": str(tz), "started_at_local": _stamp(m.started_at, tz),
                    "note": clean(m.error) or None},
        "participants": [{"label": p.label, "name": clean(p.name) or None, "is_local_user": p.is_local_user,
                          "speaking_ms": sum(s.end_ms - s.start_ms for s in m.segments if s.participant_id == p.id)}
                         for p in m.participants],
        "summary": ({"overview": clean(m.summary.overview), "key_points": [clean(k) for k in (m.summary.key_points or [])],
                     "decisions": [clean(d) for d in m.summary.decisions],
                     "minutes": [{"start_ms": x.get("start_ms"), "point": clean(x.get("point"))} for x in m.summary.minutes]}
                    if m.summary else None),
        "tasks_schema": TASKS_SCHEMA,
        "tasks": [task_dict(t, people) for t in chosen],
        "generated_at": datetime.now(UTC).isoformat(timespec="seconds"),
    }
    if transcript:
        doc["transcript"] = [{"start_ms": s.start_ms, "end_ms": s.end_ms, "source": s.source, "text": clean(s.text),
                              "speaker": (lambda p: (clean(p.name or p.label) if p else None))(people.get(s.participant_id or ""))}
                             for s in m.segments]
    return doc


def _groups(doc: dict) -> list[dict]:
    """Consecutive segments of one speaker read as one paragraph."""
    out: list[dict] = []
    for s in doc.get("transcript", []):
        if out and out[-1]["speaker"] == s["speaker"]:
            out[-1]["text"] += " " + s["text"]
        else:
            out.append({"speaker": s["speaker"], "start_ms": s["start_ms"], "text": s["text"]})
    return out


def _stamp(dt: datetime | None, tz: ZoneInfo) -> str:
    """'15-10-2026 09:30 WIB': the clock a participant would have read, with the zone named."""
    if dt is None:
        return ""
    local = to_zone(dt, tz)
    return f"{local:%d-%m-%Y %H:%M} {label(tz, local)}"


# ------------------------------------------------------------------------------------------ JSON
def to_json(doc: dict) -> bytes:
    return json.dumps(doc, ensure_ascii=False, indent=2).encode("utf-8")


# -------------------------------------------------------------------------------------- Markdown
_MD_SPECIAL = re.compile(r"([\\`*_{}\[\]<>|])")


def md(text: str) -> str:
    """Escape what would be read as markup or raw HTML: the text comes from speech recognition and an LLM."""
    text = _MD_SPECIAL.sub(r"\\\1", clean(text))
    return re.sub(r"(?m)^(\s*)([#>+\-]|\d+\.)(?=\s)", r"\1\\\2", text)


def to_markdown(doc: dict) -> bytes:
    m = doc["meeting"]
    L = [f"# {md(m['title'])}", "", f"- **Tanggal:** {m['started_at_local']}", f"- **Bahasa:** {m['language']}"]
    if doc["participants"]:
        L.append("- **Peserta:** " + ", ".join(md(p["name"] or p["label"]) for p in doc["participants"]))
    if m["note"]:
        L += ["", f"> Catatan: {md(m['note'])}"]
    s = doc["summary"]
    if s:
        L += ["", "## Ringkasan", "", md(s["overview"])]
        if s.get("key_points"):
            L += ["", "## Poin penting", ""] + [f"- {md(k)}" for k in s["key_points"]]
        if s["decisions"]:
            L += ["", "## Keputusan", ""] + [f"- {md(d)}" for d in s["decisions"]]
        if s["minutes"]:
            L += ["", "## Jalannya rapat", ""] + [f"- **{clock(x['start_ms'])}** {md(x['point'])}" for x in s["minutes"]]
    if doc["tasks"]:
        L += ["", "## Task", ""]
        for t in doc["tasks"]:
            meta = [f"PIC: {md(t['assignee'])}" if t["assignee"] else None,
                    (f"Tenggat: {t['deadline']}" + (" (perkiraan)" if t["deadline_is_approximate"] else "")) if t["deadline"]
                    else (f"Tenggat: {md(t['deadline_text'])}" if t["deadline_text"] else None),
                    f"Estimasi: {minutes_text(t['estimate_minutes'])}" if t["estimate_minutes"] else None,
                    f"Prioritas: {PRIORITY.get(t['priority'], t['priority'])}",
                    None if t["status"] == "confirmed" else STATUS.get(t["status"])]
            box = "x" if t["status"] == "confirmed" else " "
            L.append(f"- [{box}] **{md(t['title'])}** — " + " · ".join(x for x in meta if x))
            if t["description"]:
                L.append(f"  {md(t['description'])}")
            if t["evidence"]["quote"]:
                at = f" ({clock(t['evidence']['start_ms'])})" if t["evidence"]["start_ms"] is not None else ""
                L.append(f"  > {md(t['evidence']['quote'])}{at}")
    g = _groups(doc)
    if g:
        L += ["", "## Transkrip", ""]
        for x in g:
            L += [f"**[{clock(x['start_ms'])}] {md(x['speaker'] or 'Tidak dikenal')}:** {md(x['text'])}", ""]
    return ("\n".join(L).rstrip() + "\n").encode("utf-8")


# ------------------------------------------------------------------------------------------ Word
def _deep_clean(x):
    if isinstance(x, str):
        return clean(x)
    if isinstance(x, list):
        return [_deep_clean(i) for i in x]
    if isinstance(x, dict):
        return {k: _deep_clean(v) for k, v in x.items()}
    return x


def to_docx(doc: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = _deep_clean(doc)  # whatever built the document, the Word writer must never see a control character

    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    m = doc["meeting"]
    d.add_heading(clean(m["title"]) or "Rapat", level=0)
    meta = d.add_paragraph()
    meta.add_run("Tanggal: ").bold = True
    meta.add_run(m["started_at_local"] + "\n")
    if doc["participants"]:
        meta.add_run("Peserta: ").bold = True
        meta.add_run(", ".join(clean(p["name"] or p["label"]) for p in doc["participants"]))
    if m["note"]:
        n = d.add_paragraph()
        n.add_run("Catatan: ").bold = True
        n.add_run(clean(m["note"]))
    s = doc["summary"]
    if s:
        d.add_heading("Ringkasan", level=1)
        d.add_paragraph(s["overview"])
        if s.get("key_points"):
            d.add_heading("Poin penting", level=1)
            for x in s["key_points"]:
                d.add_paragraph(x, style="List Bullet")
        if s["decisions"]:
            d.add_heading("Keputusan", level=1)
            for x in s["decisions"]:
                d.add_paragraph(x, style="List Bullet")
        if s["minutes"]:
            d.add_heading("Jalannya rapat", level=1)
            for x in s["minutes"]:
                p = d.add_paragraph(style="List Bullet")
                p.add_run(clock(x["start_ms"]) + "  ").bold = True
                p.add_run(x["point"])
    if doc["tasks"]:
        d.add_heading("Task", level=1)
        t = d.add_table(rows=1, cols=6)
        t.style = "Light Grid Accent 1"
        for i, h in enumerate(("Task", "PIC", "Tenggat", "Estimasi", "Prioritas", "Status")):
            t.rows[0].cells[i].text = h
        for x in doc["tasks"]:
            c = t.add_row().cells
            c[0].text = x["title"] + (f"\n{x['description']}" if x["description"] else "")
            c[1].text = x["assignee"] or "-"
            c[2].text = (x["deadline"] + (" (perkiraan)" if x["deadline_is_approximate"] else "")) if x["deadline"] else (x["deadline_text"] or "-")
            c[3].text = minutes_text(x["estimate_minutes"]) or "-"
            c[4].text = PRIORITY.get(x["priority"], x["priority"])
            c[5].text = STATUS.get(x["status"], x["status"])
    g = _groups(doc)
    if g:
        d.add_heading("Transkrip", level=1)
        for x in g:
            p = d.add_paragraph()
            p.add_run(f"[{clock(x['start_ms'])}] {clean(x['speaker'] or 'Tidak dikenal')}: ").bold = True
            p.add_run(x["text"])
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------------------------- filenames
MEDIA = {
    "md": "text/markdown; charset=utf-8",
    "json": "application/json",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


def render(doc: dict, fmt: str) -> bytes:
    return {"md": to_markdown, "json": to_json, "docx": to_docx}[fmt](doc)


def filename(title: str, started: date | datetime | None, fmt: str) -> str:
    base = unicodedata.normalize("NFKD", title or "rapat").encode("ascii", "ignore").decode()
    base = re.sub(r"[^A-Za-z0-9]+", "-", base).strip("-").lower()[:60] or "rapat"
    day = started.strftime("%Y%m%d") if started else ""
    return f"rekamind-{day + '-' if day else ''}{base}.{fmt}"
