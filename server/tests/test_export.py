import io

from test_trash import _processed_meeting

from rekamind_server import export


async def _confirm_first(client, m):
    tasks = (await client.get(f"/api/meetings/{m}")).json()["tasks"]
    r = await client.patch(f"/api/tasks/{tasks[0]['id']}", json={"status": "confirmed"})
    assert r.status_code == 200
    return tasks


async def test_markdown_export_has_everything_and_only_confirmed_tasks_by_default(client):
    m = await _processed_meeting(client)
    tasks = await _confirm_first(client, m)
    r = await client.get(f"/api/meetings/{m}/export", params={"format": "md"})
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/markdown")
    assert 'filename="rekamind-' in r.headers["content-disposition"] and r.headers["content-disposition"].endswith('.md"')
    text = r.text
    for part in ("# Akan dihapus", "## Ringkasan", "## Keputusan", "## Jalannya rapat", "## Task", "## Transkrip", "Budi tolong siapkan draft kontrak"):
        assert part in text, part
    assert text.count("- [x]") == 1 and "- [ ]" not in text, "unconfirmed suggestions are not exported by default"
    full = (await client.get(f"/api/meetings/{m}/export", params={"format": "md", "tasks": "all"})).text
    assert full.count("- [x]") == 1 and full.count("- [ ]") == len({t["title"] for t in tasks}) - 1
    assert "## Transkrip" not in (await client.get(f"/api/meetings/{m}/export", params={"format": "md", "transcript": "false"})).text


async def test_json_export_is_the_versioned_contract(client):
    m = await _processed_meeting(client)
    await _confirm_first(client, m)
    d = (await client.get(f"/api/meetings/{m}/export", params={"format": "json"})).json()
    assert d["schema"] == "rekamind.meeting/v1" and d["tasks_schema"] == "rekamind.tasks/v1"
    assert d["meeting"]["id"] == m and d["summary"]["overview"] and len(d["tasks"]) == 1
    t = d["tasks"][0]
    assert set(t) >= {"id", "title", "assignee", "deadline", "deadline_is_approximate", "estimate_minutes", "priority", "evidence", "status"}
    assert len(d["transcript"]) == 4 and {p["label"] for p in d["participants"]} >= {"ME"}


async def test_word_export_is_a_real_docx_with_the_content_in_it(client):
    from docx import Document

    m = await _processed_meeting(client)
    await _confirm_first(client, m)
    r = await client.get(f"/api/meetings/{m}/export")  # docx is the default
    assert r.headers["content-type"] == export.MEDIA["docx"]
    doc = Document(io.BytesIO(r.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Akan dihapus" in text and "Ringkasan" in text and "Transkrip" in text and "Budi tolong siapkan draft kontrak" in text
    assert len(doc.tables) == 1 and doc.tables[0].rows[0].cells[0].text == "Task" and len(doc.tables[0].rows) == 2


def test_hostile_text_cannot_inject_markup_or_break_the_word_file():
    nasty = "<script>alert(1)</script> [x](javascript:alert(1)) **bold** \x00\x08 `code`\n# heading\n> quote"
    out = export.md(nasty)
    assert "<script>" not in out and "\\<script\\>" in out and "\\[x\\]" in out and "\x00" not in out
    assert "\n\\# heading" in out and "\n\\> quote" in out
    from docx import Document

    doc = {"meeting": {"title": nasty, "started_at": None, "language": "id", "note": nasty, "id": "x", "started_at_local": nasty, "ended_at": None, "status": "review"},
           "participants": [], "summary": {"overview": nasty, "decisions": [nasty], "minutes": []}, "tasks": [],
           "transcript": [{"start_ms": 0, "end_ms": 1, "source": "mic", "text": nasty, "speaker": nasty}]}
    Document(io.BytesIO(export.to_docx(doc)))  # must not raise on control characters


async def test_filenames_are_safe_and_the_export_follows_ownership(client):
    assert export.filename('../../Rapat "Q3" <vendor>\r\nX', None, "md") == "rekamind-rapat-q3-vendor-x.md"
    assert export.filename("日本語", None, "docx") == "rekamind-rapat.docx"
    m = await _processed_meeting(client)
    assert (await client.get("/api/meetings/" + "0" * 32 + "/export")).status_code == 404
    assert (await client.get(f"/api/meetings/{m}/export", params={"format": "pdf"})).status_code == 422


async def test_an_unfinished_meeting_cannot_be_exported(client):
    m = (await client.post("/api/meetings", json={"title": "x"})).json()["id"]
    assert (await client.get(f"/api/meetings/{m}/export")).status_code == 409
